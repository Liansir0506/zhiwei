"""Three-format ingestion -> stored anchor -> fresh scoped evidence retrieval."""
import uuid
import csv
import copy
import pytest
from unittest.mock import Mock
from sqlalchemy import text
from pymilvus import MilvusClient
from qa_core.config.settings import get_settings
from qa_core.storage.kb_dao import KbDao
from qa_core.storage.location_dao import LocationDao
from qa_core.retrieval.searcher import Searcher
from qa_core.retrieval.filters import _to_doc
from qa_core.schemas.evidence import evidence_of
from qa_core.application.evidence_service import EvidenceService
from qa_core.security.policy import AccessScope
from scripts.init_collections import build
from scripts import rebuild_kb


def test_three_format_ingestion_persists_content_bound_anchors(tmp_path, monkeypatch):
    import pymupdf
    from docx import Document
    s = get_settings(); token = uuid.uuid4().hex[:12]
    collection, version = f"test_location_{token}", f"location-it-{token}"
    s = s.model_copy(update={"doc_collection": collection})
    monkeypatch.setattr(rebuild_kb, "get_settings", lambda: s)
    client = MilvusClient(uri=s.milvus_uri, timeout=15); dao = KbDao(s.mysql_dsn)
    paths = [tmp_path / "manual_anchor.md", tmp_path / "manual_anchor.pdf", tmp_path / "manual_anchor.docx"]
    paths[0].write_text("# MD Anchor\nUnique Markdown evidence.", encoding="utf-8")
    with pymupdf.open() as d:
        d.new_page().insert_text((72, 72), "# PDF Anchor\nUnique PDF evidence."); d.save(paths[1])
    d = Document(); d.add_paragraph("# DOCX Anchor"); d.add_paragraph("Unique DOCX evidence."); d.save(paths[2])
    embedder = Mock(); embedder.embed_passages.side_effect = lambda texts: [[1.0] + [0.0] * 1023 for _ in texts]
    from types import SimpleNamespace
    from transformers import AutoTokenizer
    embedder.model = SimpleNamespace(tokenizer=AutoTokenizer.from_pretrained(s.embedding_model_path, local_files_only=True),
                                     max_seq_length=8192)
    identities = []
    try:
        build(client, collection, False); dao.ensure_version(version)
        stats = {"done": 0, "skipped": 0, "failed": 0, "duplicates": 0}
        for i, path in enumerate(paths):
            rebuild_kb.process_file(client, dao, embedder, path, version, set(), stats,
                strict_metadata=True, record_locations=True, tokenizer_chunking=True, extra_meta={"asset_id": f"anchor-{token}-{i}",
                    "tenant_id": "default", "visibility": "public", "security_level": "internal", "status": "active"})
        assert stats["done"] == 3 and stats["failed"] == 0
        client.flush(collection)
        searcher = Searcher(client=client, embedder=Mock(), cache=Mock(), settings=s)
        scope = AccessScope.build(roles=["ops"])
        rows = client.query(collection_name=collection, filter=f'kb_version == "{version}"',
                            output_fields=["*"], limit=10, consistency_level="Strong")
        assert len(rows) == 3
        service = EvidenceService(searcher, Mock(), s)
        for row in rows:
            doc = _to_doc({"id": row["id"], "entity": row}); identities.append(LocationDao.identity(doc))
            e = service.source(evidence_of(doc)["evidence_id"], scope)
            location = e["location"]; assert location["kind"] == "source_spans"
            assert location["file_sha256"] == rebuild_kb.file_hash_of(next(p for p in paths if p.name == doc.source))
            expected = {"md": "lines", "pdf": "page", "docx": "paragraph"}[location["format"]]
            assert location["matches"][0][0]["kind"] == expected
            assert service.source(e["evidence_id"], AccessScope.build(tenant_id="other", roles=["admin"])) is None
        assert dao.get_version(version)["status"] == "STAGED"
    finally:
        if client.has_collection(collection): client.drop_collection(collection)
        client.close()
        with dao.engine.begin() as conn:
            for identity in identities:
                conn.execute(text("DELETE FROM evidence_locations WHERE identity_hash=:id"), {"id": identity})
            conn.execute(text("DELETE FROM ingest_records WHERE kb_version=:v"), {"v": version})
            conn.execute(text("DELETE FROM kb_versions WHERE version=:v"), {"v": version})


@pytest.mark.parametrize("kind", ["faq", "fault_case"])
def test_csv_ingestion_binds_record_location_and_faq_answer(tmp_path, monkeypatch, kind):
    token = uuid.uuid4().hex[:12]
    collection, version = f"test_csv_location_{token}", f"csv-location-{token}"
    original = get_settings()
    settings = original.model_copy(update={"faq_collection" if kind == "faq" else "doc_collection": collection})
    monkeypatch.setattr(rebuild_kb, "get_settings", lambda: settings)
    path = tmp_path / (kind + "_anchor.csv")
    source = "fixture-record-" + token
    row = ({"source": source, "question": "模拟问题？", "answer": "模拟答案\n第二行"} if kind == "faq" else
           {"ticket_id": source, "title": "模拟故障", "phenomenon": "仅测试现象", "root_cause": "仅测试原因", "resolution": "仅测试处理"})
    row["network_tech"] = "4G"
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(row)); writer.writeheader(); writer.writerow(row)
    client = MilvusClient(uri=settings.milvus_uri, timeout=20)
    dao = KbDao(settings.mysql_dsn)
    identities = []
    try:
        build(client, collection, kind == "faq"); dao.ensure_version(version)
        embedder = Mock(); embedder.embed_passages.side_effect = lambda texts: [[1.0] + [0.0] * 1023 for _ in texts]
        stats = {"done": 0, "skipped": 0, "failed": 0, "duplicates": 0}
        rebuild_kb.process_file(client, dao, embedder, path, version, set(), stats,
            strict_metadata=True, record_locations=True,
            extra_meta={"asset_id": "csv-fixture-" + token, "source_id": source, "source_system": "test_fixture",
                "source_version": "fixture-r1", "tenant_id": "default", "visibility": "public",
                "security_level": "internal", "status": "active", "network_tech": "5G"})
        assert stats["done"] == 1 and stats["failed"] == 0
        client.flush(collection)
        stored = client.query(collection_name=collection, filter=f'kb_version == "{version}"',
                              output_fields=["*"], limit=2, consistency_level="Strong")
        assert len(stored) == 1
        assert stored[0]["network_tech"] == "4G"
        doc = _to_doc({"id": stored[0]["id"], "entity": stored[0]})
        identities.append(LocationDao.identity(doc))
        locator = LocationDao(dao.engine).get(doc)
        assert locator["file_sha256"] == rebuild_kb.file_hash_of(path)
        assert locator["matches"][0][0]["record"] == 1
        assert locator["matches"][0][0]["line_end"] == (3 if kind == "faq" else 2)
        searcher = Searcher(client=client, settings=settings, embedder=Mock(), cache=Mock())
        service = EvidenceService(searcher, Mock(), settings)
        scope = AccessScope.build(roles=["ops"])
        evidence = evidence_of(doc, "faq" if kind == "faq" else "doc")
        reread = service.source(evidence["evidence_id"], scope)
        assert reread["location"] == locator
        assert reread["text"] == (doc.answer if kind == "faq" else doc.text)
        if kind == "faq":
            changed = copy.deepcopy(doc); changed.answer += "invented addition"
            assert LocationDao(dao.engine).get(changed) is None
            client.upsert(collection_name=collection, data=[{**stored[0], "answer": changed.answer}])
            client.flush(collection)
            assert service.source(evidence["evidence_id"], scope) is None
        assert service.source(evidence["evidence_id"], AccessScope.build(tenant_id="other", roles=["admin"])) is None
        # Metadata-only upsert must move the partition key without leaving two PKs.
        moved = {**stored[0], "network_tech": "5G"}
        client.upsert(collection_name=collection, data=[moved]); client.flush(collection)
        moved_rows = client.query(collection_name=collection, filter=f'id == "{doc.id}"',
                                  output_fields=["id", "network_tech", "dense"], limit=3,
                                  consistency_level="Strong")
        assert len(moved_rows) == 1 and moved_rows[0]["network_tech"] == "5G"
        assert moved_rows[0]["dense"] == stored[0]["dense"]
        assert dao.get_version(version)["status"] == "STAGED"
    finally:
        if client.has_collection(collection): client.drop_collection(collection)
        client.close()
        with dao.engine.begin() as conn:
            for identity in identities:
                conn.execute(text("DELETE FROM evidence_locations WHERE identity_hash=:id"), {"id": identity})
            conn.execute(text("DELETE FROM ingest_records WHERE kb_version=:v"), {"v": version})
            conn.execute(text("DELETE FROM kb_version_manifests WHERE version=:v"), {"v": version})
            conn.execute(text("DELETE FROM kb_versions WHERE version=:v AND status='STAGED'"), {"v": version})
        dao.engine.dispose()
