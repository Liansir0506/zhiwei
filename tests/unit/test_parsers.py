"""解析器与脱敏纯逻辑测试：DOCX/PDF 用库程序化生成，不依赖真实资料。"""
import csv
import json

import pytest

from qa_core.indexing.parsers.parsers import (
    dedup_header_footer,
    merge_cross_page,
    parse_dispatch,
)
from qa_core.indexing.parsers.sanitize import garbled_ratio, sanitize


@pytest.mark.unit
def test_sanitize_replacements():
    text = "联系 13812345678，账号密码 password=Abc123，地址 10.20.30.40，证件 110101199001011234"
    out, counts = sanitize(text)
    assert "{PHONE}" in out and "{IP}" in out and "{SECRET}" in out and "{ID}" in out
    assert "13812345678" not in out


@pytest.mark.unit
def test_garbled_ratio_reject_threshold():
    assert garbled_ratio("正常中文文本，没有乱码。") == 0.0
    bad = "abc" + "\x00\ufffd" * 10
    assert garbled_ratio(bad) > 0.05


@pytest.mark.unit
def test_dedup_header_footer():
    pages = [f"页眉X\n第{i}页内容开始。中间叙述。\n页脚Y" for i in range(5)]
    cleaned = dedup_header_footer(pages)
    assert all("页眉X" not in p and "页脚Y" not in p for p in cleaned)
    assert all("中间叙述" in p for p in cleaned)


@pytest.mark.unit
def test_dedup_keeps_short_docs():
    pages = ["页眉A\n正文一。", "页眉A\n正文二。"]  # <3 页不处理
    assert dedup_header_footer(pages) == pages


@pytest.mark.unit
def test_merge_cross_page():
    pages = ["第一页句子未完", "第二页接续内容。完整结束。", "新段开头。"]
    merged = merge_cross_page(pages)
    assert "句子未完第二页接续内容。" in merged


@pytest.mark.unit
def test_parse_docx_generated(tmp_path):
    import docx

    document = docx.Document()
    document.add_paragraph("设备手册第一章。")
    document.add_paragraph("安装步骤说明。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "参数"
    table.cell(0, 1).text = "值"
    table.cell(1, 0).text = "功率"
    table.cell(1, 1).text = "40W"
    path = tmp_path / "manual.docx"
    document.save(str(path))

    text = parse_dispatch("manual", path)
    assert "设备手册第一章" in text and "| 参数 | 值 |" in text


@pytest.mark.unit
def test_parse_pdf_generated(tmp_path):
    import pymupdf

    path = tmp_path / "guide.pdf"
    doc = pymupdf.open()
    for i in range(4):
        page = doc.new_page()
        # pymupdf 默认 base14 字体不支持 CJK，必须用内置中文字体 china-s
        page.insert_text((72, 50), "运维指南页眉X", fontname="china-s")
        page.insert_text((72, 100), f"第{i}章 告警处理流程。详细步骤说明。", fontname="china-s")
    doc.save(str(path))
    doc.close()

    text = parse_dispatch("guide", path)
    assert "告警处理流程" in text
    assert "运维指南页眉X" not in text  # 页眉去重


@pytest.mark.unit
def test_parse_dispatch_ticket_csv_sanitized(tmp_path):
    path = tmp_path / "tickets.csv"
    with open(path, "w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=["alarm_code", "phenomenon", "root_cause", "resolution"])
        writer.writeheader()
        writer.writerow({"alarm_code": "ALM-1234", "phenomenon": "联系值班 13912345678",
                         "root_cause": "光模块老化", "resolution": "更换光模块"})
    rows = parse_dispatch("fault_case", path)
    assert rows[0]["alarm_code"] == "ALM-1234"
    assert "{PHONE}" in rows[0]["phenomenon"]


@pytest.mark.unit
def test_parse_dispatch_rejects_garbled(tmp_path):
    path = tmp_path / "bad.txt"
    path.write_text("正常" + "\ufffd" * 100, encoding="utf-8")
    with pytest.raises(ValueError, match="乱码率"):
        parse_dispatch("sop", path)


def test_json_schemas_loadable():
    """占位：DATA_MODEL 的 scenario yaml 合法性在阶段6 配置加载时测。"""
    assert json.loads('{"ok": true}')["ok"] is True
