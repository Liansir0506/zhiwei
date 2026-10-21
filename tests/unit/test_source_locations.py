import copy
import shutil
import subprocess
from pathlib import Path
import pytest
from qa_core.indexing.parsers.locations import locate_chunks
from qa_core.indexing.parsers.parsers import parse_dispatch, parse_pdf
from qa_core.indexing.chunkers.registry import chunk_manual, chunk_sop
from qa_core.storage.location_dao import LocationDao
from qa_core.schemas.retrieval import RetrievedDoc


def test_markdown_lines_and_sop_context(tmp_path):
    p = tmp_path / "sop_demo.md"
    p.write_text("# AMF 恢复流程\n适用：沙箱\n步骤 1：只读检查。\n步骤 2：审批后回滚。\n", encoding="utf-8")
    positions = locate_chunks(p, chunk_sop(parse_dispatch("sop", p)))
    assert positions[1]["context_prefix"]
    assert positions[1]["matches"] == [[{"kind": "lines", "line_start": 3, "line_end": 3}]]
    assert positions[2]["matches"][0][0]["line_start"] == 4


def test_duplicate_text_is_ambiguous_not_guessed(tmp_path):
    p = tmp_path / "manual_demo.md"; p.write_text("# 一\n重复说明。\n# 二\n重复说明。", encoding="utf-8")
    positions = locate_chunks(p, chunk_manual(parse_dispatch("manual", p)))
    assert all(v["ambiguous"] and len(v["matches"]) == 2 for v in positions)


def test_pdf_page_anchor_and_empty_scan_rejection(tmp_path):
    import pymupdf
    p = tmp_path / "manual_demo.pdf"
    with pymupdf.open() as d:
        d.new_page().insert_text((72, 72), "# Section One\nFirst page evidence.")
        d.new_page().insert_text((72, 72), "# Section Two\nSecond page evidence.")
        d.save(p)
    positions = locate_chunks(p, chunk_manual(parse_dispatch("manual", p)))
    assert [v["matches"][0][0]["page"] for v in positions] == [1, 2]
    empty = tmp_path / "scan.pdf"
    with pymupdf.open() as d: d.new_page(); d.save(empty)
    with pytest.raises(ValueError): parse_pdf(empty)


def test_docx_blank_paragraph_count_preserved(tmp_path):
    from docx import Document
    p = tmp_path / "manual_demo.docx"
    d = Document(); d.add_paragraph("# Section One"); d.add_paragraph(""); d.add_paragraph("Paragraph evidence."); d.save(p)
    positions = locate_chunks(p, chunk_manual(parse_dispatch("manual", p)))
    assert positions[0]["matches"] == [[{"kind": "paragraph", "paragraph": 3}]]


def test_location_identity_bound_to_tenant_version_and_text():
    doc = RetrievedDoc("c", "evidence", .8, tenant_id="t1", asset_id="a1", kb_version="v1", source_version="r1")
    before = LocationDao.identity(doc)
    for field in ("tenant_id", "asset_id", "kb_version", "source_version", "text", "id"):
        other = copy.deepcopy(doc); setattr(other, field, "changed")
        assert LocationDao.identity(other) != before


def test_faq_location_identity_changes_when_only_the_answer_changes():
    doc = RetrievedDoc("c", "question", .8, doc_type="faq", answer="first answer",
                       tenant_id="t1", asset_id="a1", kb_version="v1", source_version="r1")
    first = LocationDao.identity(doc)
    doc.answer = "different answer"
    assert LocationDao.identity(doc) != first


@pytest.mark.parametrize("kind", ["upload", "rebuild"])
def test_worker_records_upload_locations_and_bulk_opt_in(tmp_path, monkeypatch, kind):
    from contextlib import nullcontext
    from unittest.mock import Mock
    from types import SimpleNamespace
    from scripts import ingest_worker as worker
    monkeypatch.setattr(worker, "candidate_write", lambda *a, **kw: nullcontext())
    staged = tmp_path / "staged.md"; staged.write_text("# new source", encoding="utf-8")
    monkeypatch.setattr(worker, "get_settings", lambda: SimpleNamespace(data_dir=str(tmp_path), mysql_dsn="unused", redis_url="unused"))
    monkeypatch.setattr(worker, "CacheLayer", lambda url: Mock())
    kb = Mock(); kb.get_records_by_filename.return_value = []
    monkeypatch.setattr(worker, "KbDao", lambda dsn: kb)
    runner = Mock(return_value={"done": 1, "failed": 0})
    monkeypatch.setattr(worker, "run_rebuild", runner)
    dao = Mock(); dao.cancel_requested.return_value = False
    job = {"job_id": "job-1", "job_type": kind, "kb_version": "v-new", "mode": "incremental",
        "payload_json": {"staging_path": str(staged), "filename": "manual_new.md", "source_id": "test-source",
                         "record_locations": True, "run_quality_gate": False}}
    from qa_core.ingest.binding import execution_binding
    job["payload_json"]["execution_binding"] = execution_binding(worker.get_settings(), kind, job["payload_json"])
    worker.execute_job(job, dao, "worker")
    assert runner.call_args.kwargs["record_locations"] is True


def test_frontend_readonly_auth_and_no_admin_fallback():
    node = shutil.which("node")
    if not node: pytest.skip("Node is required for frontend contract test")
    viewer = Path(__file__).resolve().parents[2] / "static/evidence-viewer.js"
    script = """
const fs=require('fs'), vm=require('vm'), assert=require('assert');
let seen=[];
const context={window:{}, setTimeout:cb=>cb(), fetch:async(url,options)=>{seen.push({url,options});return {ok:true,json:async()=>({text:'<script>unsafe</script>'})}}};
vm.runInNewContext(fs.readFileSync(process.argv[1],'utf8'),context);
(async()=>{
 const v=context.window.EvidenceViewer;
 await v.load({evidence_id:'abc_123',source_url:'/admin/documents/chunks'}, {'X-User-Roles':'ops'});
 assert.equal(seen[0].url,'/kb/evidence/abc_123');
 assert.equal(seen[0].options.headers['X-User-Roles'],'ops');
 assert.equal(seen[0].options.cache,'no-store');
 await assert.rejects(()=>v.load({doc_id:'old'},{}));
 assert.equal(seen.length,1);
 const group=await v.loadAll({workflow_evidence:true,child_evidence:[{evidence_id:'one'},{evidence_id:'two'}]}, {'X-User-Roles':'ops'});
 assert.equal(group.length,2);
 assert.equal(seen[1].url,'/kb/evidence/one');assert.equal(seen[2].url,'/kb/evidence/two');
 assert.ok(v.location({kind:'source_spans',matches:[[{kind:'page',page:2}]]}).includes('2'));
 assert.equal(v.location({kind:'source_spans',matches:[[{kind:'csv_record',record:3,line_start:7,line_end:9}]]}), '原文位置：第 3 条记录（第 7–9 行）');
 context.fetch=async()=>({ok:false,status:404});
 await assert.rejects(()=>v.load({evidence_id:'abc'},{}),/失效/);
 await assert.rejects(()=>v.loadAll({workflow_evidence:true,child_evidence:[{evidence_id:'one'}]},{}),/失效/);
})().catch(e=>{console.error(e);process.exit(1)});
"""
    subprocess.run([node, "-e", script, str(viewer)], check=True, capture_output=True, text=True)
    html = viewer.with_name("index.html").read_text(encoding="utf-8")
    block = html.split("function addMeta(", 1)[1].split("async function jumpToDoc", 1)[0]
    assert "/admin/documents/chunks" not in block
    assert "EvidenceViewer.loadAll" in block and "box.textContent" in block
