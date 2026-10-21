"""格式解析与版式清洗（设计大纲 2.1 解析器矩阵）。

统一入口 parse_dispatch(doc_type, path)：
  fault_case -> list[dict]（表格行）
  faq        -> list[dict]
  manual / standard / guide / sop -> str（全文）
"""
from __future__ import annotations

import csv
from pathlib import Path

TERMINALS = tuple("。！？.!?:；;\n")
_GARBLED_RE = None  # 延迟到 sanitize 模块


# ---------------- 版式清洗 ----------------

def dedup_header_footer(pages: list[str]) -> list[str]:
    """移除出现在 >80% 页面首/尾的重复行（页眉页脚）。"""
    def freq(lines_getter):
        from collections import Counter

        counter: Counter[str] = Counter()
        for page in pages:
            lines = lines_getter(page)
            for ln in set(lines):
                if ln.strip():
                    counter[ln.strip()] += 1
        return counter

    n = len(pages)
    if n < 3:
        return pages
    head_freq = freq(lambda p: p.strip().splitlines()[:1])
    tail_freq = freq(lambda p: p.strip().splitlines()[-1:])
    bad_head = {ln for ln, c in head_freq.items() if c / n > 0.8}
    bad_tail = {ln for ln, c in tail_freq.items() if c / n > 0.8}

    cleaned = []
    for page in pages:
        lines = page.strip().splitlines()
        if lines and lines[0].strip() in bad_head:
            lines = lines[1:]
        if lines and lines[-1].strip() in bad_tail:
            lines = lines[:-1]
        cleaned.append("\n".join(lines))
    return cleaned


def merge_cross_page(pages: list[str]) -> str:
    """页尾句未结束（无句末标点）→ 与下页首段无缝拼接。"""
    out = ""
    for page in pages:
        page = page.strip()
        if not page:
            continue
        if out and not out.rstrip().endswith(TERMINALS):
            out += page
        else:
            out += ("\n" + page if out else page)
    return out


# ---------------- 各格式解析 ----------------

def parse_pdf(path: str | Path) -> str:
    import pymupdf

    pages: list[str] = []
    with pymupdf.open(str(path)) as doc:
        for page in doc:
            pages.append(page.get_text("text"))
    if not pages or not any(page.strip() for page in pages):
        raise ValueError(f"{path}: 无文本层（扫描件需 OCR，当前版本拒绝入库并标 quality=low 分支）")
    return merge_cross_page(dedup_header_footer(pages))


def _table_to_markdown(rows: list[list[str]]) -> str:
    if not rows:
        return ""
    header = rows[0]
    lines = ["| " + " | ".join(str(c) for c in header) + " |",
             "|" + "---|" * len(header)]
    for row in rows[1:]:
        lines.append("| " + " | ".join(str(c) for c in row) + " |")
    return "\n".join(lines)


def parse_docx(path: str | Path) -> str:
    import docx

    document = docx.Document(str(path))
    parts: list[str] = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        md = _table_to_markdown(rows)
        if md:
            parts.append(md)
    return "\n".join(parts)


def _read_rows(path: str | Path) -> list[dict]:
    """CSV / XLSX 表格行 → list[dict]（工单、FAQ）。"""
    path = str(path)
    if path.lower().endswith(".csv"):
        with open(path, encoding="utf-8-sig", newline="") as f:
            return [dict(r) for r in csv.DictReader(f)]
    import pandas as pd

    frame = pd.read_excel(path, dtype=str).fillna("")
    return frame.to_dict("records")


def parse_dispatch(doc_type: str, path: str | Path):
    """按知识源类型选择解析器。返回 str 或 list[dict]。"""
    path = Path(path)
    if doc_type in ("fault_case", "faq"):
        rows = _read_rows(path)
        # 表格类逐行脱敏
        from qa_core.indexing.parsers.sanitize import sanitize

        cleaned = []
        for row in rows:
            cleaned.append({k: sanitize(str(v))[0] for k, v in row.items()})
        return cleaned
    if path.suffix.lower() == ".pdf":
        text = parse_pdf(path)
    elif path.suffix.lower() in (".docx", ".doc"):
        text = parse_docx(path)
    elif path.suffix.lower() in (".md", ".markdown", ".txt"):
        text = Path(path).read_text(encoding="utf-8")
    else:
        raise ValueError(f"不支持的扩展名: {path.suffix} (doc_type={doc_type})")

    from qa_core.indexing.parsers.sanitize import garbled_ratio, sanitize

    ratio = garbled_ratio(text)
    if ratio > 0.05:
        raise ValueError(f"{path.name}: 乱码率 {ratio:.2%} > 5%，整份拒绝")
    return sanitize(text)[0]
